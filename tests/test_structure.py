import hashlib

from langchain_core.documents import Document

from fieldguide.ingest import chunk_documents, read_documents


def test_sections_do_not_mix_and_headings_remain_on_their_own_line():
    document = Document(
        page_content=(
            "# Survey training\nIntroductory instructions.\n"
            "## Vegetation\nRecord cover.\n- Identify plants.\n- Record NA.\n"
            "## Water\nRecord the water level."
        ), metadata={"source": "training.md", "page": None},
    )
    chunks = chunk_documents([document])
    assert len(chunks) == 3
    assert chunks[1].page_content.startswith("## Vegetation\nRecord cover.")
    assert "\n- Identify plants.\n- Record NA." in chunks[1].page_content
    assert "Water" not in chunks[1].page_content
    assert chunks[1].metadata["section"] == "Survey training > Vegetation"
    assert chunks[2].metadata["section"] == "Survey training > Water"
    assert all(chunk.metadata["title"] == "Survey training" for chunk in chunks)


def test_section_metadata_survives_chunk_splitting_and_offsets_are_exact():
    text = "## Vegetation\n" + "Record every species and its percent cover. " * 15
    chunks = chunk_documents([
        Document(page_content=text, metadata={"source": "training.md", "page": None}),
    ], size=160, overlap=45)
    assert len(chunks) > 1
    assert all(chunk.metadata["section"] == "Vegetation" for chunk in chunks)
    for chunk in chunks:
        assert len(chunk.page_content) <= 160
        assert text[chunk.metadata["start_index"]:chunk.metadata["end_index"]] == chunk.page_content


def test_pdf_page_spans_and_explicit_heading_boundaries():
    documents = [Document(page_content=text, metadata={"source": "sop.pdf", "page": i})
                 for i, text in enumerate([
                     "## Vegetation\nRecord the species and",
                     "its cover.\n## Water\nRecord water level.",
                 ], 1)]
    chunks = chunk_documents(documents)
    assert len(chunks) == 2
    assert "species and its cover." in chunks[0].page_content
    assert (chunks[0].metadata["page"], chunks[0].metadata["page_end"]) == (1, 2)
    assert chunks[1].metadata["page"] == 2
    assert chunks[1].metadata["section"] == "Water"


def test_read_records_file_hash_and_docx_heading_styles(tmp_path):
    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_heading("Field methods", level=1)
    doc.add_paragraph("Record the station.")
    doc.add_heading("Review", level=2)
    doc.add_paragraph("Check the notebook.")
    path = tmp_path / "training.docx"
    doc.save(path)
    documents, warnings = read_documents(tmp_path)
    assert not warnings
    assert documents[0].page_content.startswith("# Field methods\n\nRecord the station.")
    assert documents[0].metadata["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    chunks = chunk_documents(documents)
    assert chunks[-1].metadata["section"] == "Field methods > Review"
    assert chunks[-1].metadata["source_sha256"] == documents[0].metadata["source_sha256"]


def test_legacy_sentence_strategy_is_available():
    chunks = chunk_documents([
        Document(page_content="Record species. Record cover.",
                 metadata={"source": "training.txt", "page": None}),
    ], strategy="sentence")
    assert chunks[0].metadata["chunk_method"] == "sentence-v1"
    assert chunks[0].metadata["end_index"] == len(chunks[0].page_content)


def test_markdown_table_rows_remain_separate():
    text = "| Site | Reading |\n| --- | --- |\n| A | 10 |\n| B | 20 |"
    chunks = chunk_documents([
        Document(page_content=text, metadata={"source": "table.md", "page": None}),
    ])
    assert chunks[0].page_content == text


def test_wrapped_list_items_stay_with_their_continuations():
    text = "- Record the species and\n  percent cover.\n- Record the plot and\n  observer initials."
    chunks = chunk_documents([
        Document(page_content=text, metadata={"source": "list.md", "page": None}),
    ])
    assert chunks[0].page_content == (
        "- Record the species and percent cover.\n- Record the plot and observer initials."
    )


def test_decimal_sections_do_not_mix_with_previous_procedure():
    text = (
        "2.1 Marker observations\n1. Locate the marker.\n2. Record its condition.\n"
        "2.2 Elevation measurements\nRead the instrument before and after the survey."
    )
    chunks = chunk_documents([
        Document(page_content=text, metadata={"source": "survey.pdf", "page": 1}),
    ], size=1000)
    assert len(chunks) == 2
    assert "Marker observations" in chunks[0].page_content
    assert "2. Record its condition." in chunks[0].page_content
    assert chunks[1].page_content.startswith("2.2 Elevation measurements\n")
    assert "marker" not in chunks[1].page_content
    assert chunks[1].metadata["section"] == "Elevation measurements"


def test_page_leading_decimal_heading_is_separate_and_retains_full_text():
    heading = "2.2 Elevation " + "measurement instructions " * 8
    documents = [
        Document(page_content="2.1 Marker observations\nRead the marker.",
                 metadata={"source": "survey.pdf", "page": 1}),
        Document(page_content=heading + "\nRead before and after the survey.",
                 metadata={"source": "survey.pdf", "page": 2}),
    ]
    chunks = chunk_documents(documents)
    assert len(chunks) == 2
    assert chunks[1].page_content.startswith(heading.strip())
    assert chunks[1].metadata["page"] == 2
    assert chunks[1].metadata["page_end"] == 2
    assert len(chunks[1].metadata["section"]) == 120

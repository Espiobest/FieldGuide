import json

import numpy as np
import pytest
from langchain_core.documents import Document

from fieldguide.ingest import chunk_documents, read_documents
from fieldguide.store import LocalIndex


class TinyEmbeddings:
    def embed_documents(self, texts):
        return [self.embed_query(t) for t in texts]

    def embed_query(self, text):
        return [1.0 if "plant" in text else 0.0, 1.0 if "water" in text else 0.0]


def test_ingestion_and_stable_source_locations(tmp_path):
    (tmp_path / "notes.md").write_text("Plant observation. " * 30, encoding="utf-8")
    (tmp_path / ".secret.txt").write_text("hidden", encoding="utf-8")
    documents, warnings = read_documents(tmp_path)
    first = chunk_documents(documents, size=100, overlap=20)
    second = chunk_documents(documents, size=100, overlap=20)
    assert not warnings
    assert len(first) > 1 and first == second
    assert all(d.metadata["source"] == "notes.md" for d in first)
    assert all(len(d.page_content) <= 100 for d in first)
    assert len({d.metadata["chunk_id"] for d in first}) == len(first)


def test_empty_and_invalid_chunk_config(tmp_path):
    with pytest.raises(ValueError, match="No readable"):
        read_documents(tmp_path)
    with pytest.raises(ValueError, match="overlap"):
        chunk_documents([], size=100, overlap=100)


def test_docx_tables_are_ingested(tmp_path):
    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_paragraph("Training notes")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Units"
    table.cell(0, 1).text = "centimeters"
    doc.add_paragraph("Record the values after checking units.")
    doc.save(tmp_path / "table.docx")
    documents, warnings = read_documents(tmp_path)
    assert "centimeters" in documents[0].page_content
    assert documents[0].page_content.index("centimeters") < documents[0].page_content.index(
        "Record the values"
    )
    assert documents[0].metadata["page"] is None and not warnings


def build_index(tmp_path, monkeypatch):
    monkeypatch.setattr("fieldguide.store.embeddings", lambda model: TinyEmbeddings())
    docs = [
        Document(page_content="plant", metadata={"source": "plants.md", "chunk_id": "a"}),
        Document(page_content="water", metadata={"source": "water.md", "chunk_id": "b"}),
    ]
    return LocalIndex.build(docs, tmp_path)


def test_faiss_roundtrip_and_source_filter(tmp_path, monkeypatch):
    build_index(tmp_path, monkeypatch)
    loaded = LocalIndex.load(tmp_path)
    results = loaded.search("plant", k=1)
    assert results[0]["source"] == "plants.md"
    assert np.isclose(results[0]["score"], 1.0)
    assert loaded.search("plant", source="water")[0]["source"] == "water.md"
    assert loaded.search("plant", source="absent") == []
    assert not list(tmp_path.glob("*.pkl"))


def test_corrupted_index_is_rejected(tmp_path, monkeypatch):
    index = build_index(tmp_path, monkeypatch)
    (tmp_path / index.manifest["documents"]).write_text("[]", encoding="utf-8")
    with pytest.raises(ValueError, match="checksum"):
        LocalIndex.load(tmp_path)


def test_manifest_cannot_read_outside_index(tmp_path, monkeypatch):
    index = build_index(tmp_path, monkeypatch)
    manifest = {**index.manifest, "documents": "../outside.json"}
    (tmp_path / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="filename"):
        LocalIndex.load(tmp_path)


def test_empty_index_build_is_rejected_before_loading_embeddings(tmp_path, monkeypatch):
    def unexpected_embeddings(model):
        pytest.fail("Empty input should not load the embedding model")

    monkeypatch.setattr("fieldguide.store.embeddings", unexpected_embeddings)
    with pytest.raises(ValueError, match="without document chunks"):
        LocalIndex.build([], tmp_path)


def test_search_requires_embeddings(tmp_path, monkeypatch):
    build_index(tmp_path, monkeypatch)
    loaded = LocalIndex.load(tmp_path, embed=False)
    with pytest.raises(ValueError, match="embeddings enabled"):
        loaded.search("plant")


def test_manifest_count_must_match_documents(tmp_path, monkeypatch):
    index = build_index(tmp_path, monkeypatch)
    manifest = {**index.manifest, "count": 999}
    (tmp_path / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="do not match"):
        LocalIndex.load(tmp_path)

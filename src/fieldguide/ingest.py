"""Read documents and preserve source locations through chunking."""

import hashlib
import re
from pathlib import Path

from langchain_core.documents import Document

SUPPORTED = {".pdf", ".docx", ".md", ".txt"}


def read_documents(folder: Path) -> tuple[list[Document], list[str]]:
    folder = folder.resolve()
    if not folder.is_dir():
        raise ValueError(f"Document folder does not exist: {folder}")
    documents, warnings = [], []
    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        if path.is_symlink() or not path.resolve().is_relative_to(folder):
            continue
        relative = path.relative_to(folder).as_posix()
        if any(part.startswith(".") for part in path.relative_to(folder).parts):
            continue
        try:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                from pypdf import PdfReader

                pages = [page.extract_text() or "" for page in PdfReader(path).pages]
            elif suffix == ".docx":
                from docx import Document as WordDocument
                from docx.table import Table

                doc = WordDocument(path)
                blocks = []
                for block in doc.iter_inner_content():
                    if isinstance(block, Table):
                        blocks.extend(" | ".join(c.text for c in row.cells) for row in block.rows)
                    else:
                        blocks.append(block.text)
                pages = ["\n".join(blocks)]
            else:
                pages = [path.read_text(encoding="utf-8-sig")]
            for page, text in enumerate(pages, 1):
                text = text.replace("\x00", "").strip()
                if not text:
                    warnings.append(f"{relative}, page {page}: no text; may need OCR")
                    continue
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": relative,
                            "page": page if suffix == ".pdf" else None,
                        },
                    )
                )
        except Exception as exc:
            warnings.append(f"{relative}: could not read ({type(exc).__name__})")
    if not documents:
        raise ValueError("No readable documents found. Use text PDFs, DOCX, Markdown, or TXT.")
    return documents, warnings


def chunk_documents(
    documents: list[Document], size: int = 1000, overlap: int = 150,
    strategy: str = "sentence",
) -> list[Document]:
    """Use character budgets and whole-sentence overlap, except for oversized sentences."""
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    if size < 100 or overlap < 0 or overlap >= size:
        raise ValueError("Chunk size must be >= 100 and 0 <= overlap < size.")
    if strategy == "recursive":
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=size, chunk_overlap=overlap, add_start_index=True,
        )
        chunks = splitter.split_documents(documents)
    elif strategy == "sentence":
        chunks = _sentence_chunks(documents, size, overlap)
    else:
        raise ValueError("Chunk strategy must be sentence or recursive.")
    for chunk in chunks:
        identity = (
            f"{chunk.metadata['source']}:{chunk.metadata['page']}:"
            f"{chunk.metadata['start_index']}:{chunk.page_content}"
        )
        chunk.metadata["chunk_id"] = hashlib.sha256(identity.encode()).hexdigest()[:16]
    return chunks


def _normalize_layout(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[^\S\n]+", " ", text)
    # Keep explicit paragraphs, headings, and list items; unwrap PDF line breaks.
    text = re.sub(r"(?<!\n)\n(?!\n|\s*(?:#|[-*•] |\d+[.)] ))", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _document_runs(documents: list[Document]) -> list[list[Document]]:
    runs: list[list[Document]] = []
    for document in documents:
        previous = runs[-1][-1] if runs else None
        page = document.metadata.get("page")
        if (
            previous is not None and isinstance(page, int)
            and previous.metadata.get("source") == document.metadata.get("source")
            and isinstance(previous.metadata.get("page"), int)
            and page == previous.metadata["page"] + 1
        ):
            runs[-1].append(document)
        else:
            runs.append([document])
    return runs


def _units(text: str, size: int) -> list[tuple[int, int]]:
    boundaries = [0]
    for match in re.finditer(r'(?<=[.!?])["\u201d\u2019)]*\s+|\n+', text):
        prefix = text[max(0, match.start() - 100):match.start()]
        if re.search(r"(?:^|\n)\s*\d+\.$", prefix):
            continue
        if re.search(r"\b(?:e\.g|i\.e|Dr|Mr|Ms|approx|Fig|vs)\.$", prefix, re.IGNORECASE):
            continue
        boundaries.append(match.end())
    boundaries.append(len(text))
    spans = []
    for start, end in zip(boundaries, boundaries[1:], strict=False):
        while start < end and text[start].isspace():
            start += 1
        while end > start and text[end - 1].isspace():
            end -= 1
        while end - start > size:
            cut = text.rfind(" ", start, start + size + 1)
            if cut <= start:
                cut = start + size
            spans.append((start, cut))
            start = cut
            while start < end and text[start].isspace():
                start += 1
        if start < end:
            spans.append((start, end))
    return spans


def _sentence_chunks(documents: list[Document], size: int, overlap: int) -> list[Document]:
    chunks = []
    for run in _document_runs(documents):
        text, locations = "", []
        for document in run:
            normalized = _normalize_layout(document.page_content)
            if not normalized:
                continue
            if text:
                text += " "
            start = len(text)
            text += normalized
            locations.append((start, len(text), document.metadata.get("page")))
        units = _units(text, size)
        cursor = 0
        while cursor < len(units):
            end = cursor + 1
            while end < len(units) and units[end][1] - units[cursor][0] <= size:
                end += 1
            start_offset, end_offset = units[cursor][0], units[end - 1][1]
            pages = [page for start, stop, page in locations
                     if start < end_offset and stop > start_offset]
            metadata = {
                **run[0].metadata, "page": pages[0], "page_end": pages[-1],
                "start_index": start_offset, "chunk_method": "sentence-v1",
            }
            chunks.append(Document(page_content=text[start_offset:end_offset], metadata=metadata))
            if end == len(units):
                break
            next_cursor = end
            while next_cursor > cursor + 1:
                candidate = next_cursor - 1
                if end_offset - units[candidate][0] > overlap:
                    break
                if units[end][1] - units[candidate][0] > size:
                    break
                next_cursor = candidate
            cursor = next_cursor
    return chunks
